from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook

from idms_db2_phase2.domain.models import SheetMappingRow
from idms_db2_phase2.parsers.sheet_mapping.sheet_mapping_cell_utils import (
    SheetMappingCellUtils,
)
from idms_db2_phase2.parsers.sheet_mapping.sheet_mapping_diagnostics import (
    SheetMappingDiagnostics,
)
from idms_db2_phase2.parsers.sheet_mapping.sheet_mapping_row_mapper import (
    SheetMappingRowMapper,
)
from rules.sheet_mapping_parser_rules import (
    CSV_SAMPLE_LENGTH,
    CSV_TEXT_ENCODING,
    DIAG_CSV_DECODED_LEN_TEMPLATE,
    DIAG_CSV_EMPTY,
    DIAG_CSV_HEADERS_TEMPLATE,
    DIAG_CSV_NO_ROWS,
    DIAG_CSV_SAMPLE_TEMPLATE,
    DIAG_CSV_USEFUL_ROWS_TEMPLATE,
    DIAG_FILE_NAME_TEMPLATE,
    DIAG_FILE_SIZE_TEMPLATE,
    DIAG_NO_FILE,
    DIAG_SHEET_NO_HEADER_TEMPLATE,
    DIAG_SHEET_USEFUL_ROWS_TEMPLATE,
    DIAG_UNSUPPORTED_XLS,
    DIAG_XLSX_EMPTY,
    DIAG_XLSX_USEFUL_ROWS_TEMPLATE,
    LEGACY_XLS_SUFFIX,
    XLSX_SUFFIX,
)


class SheetMappingParser(
    SheetMappingCellUtils,
    SheetMappingRowMapper,
    SheetMappingDiagnostics,
):
    """Parses Sheet Mapping files from CSV/text or XLSX.

    Column names, aliases, and header detection groups live in
    catalogs/sheet_mapping_schema.py; regex in patterns/sheet_mapping_patterns.py;
    messages and magic numbers in rules/sheet_mapping_parser_rules.py.
    Cell utils, row mapping, and diagnostics are provided by mixins.
    """

    def __init__(self) -> None:
        self.diagnostics: list[str] = []

    def parse_uploaded_file(self, uploaded_file) -> list[SheetMappingRow]:
        self.diagnostics = []

        if uploaded_file is None:
            self.diagnostics.append(DIAG_NO_FILE)
            return []

        file_name = str(uploaded_file.name or "").lower()
        raw_bytes = uploaded_file.getvalue()

        self.diagnostics.append(DIAG_FILE_NAME_TEMPLATE.format(name=file_name))
        self.diagnostics.append(DIAG_FILE_SIZE_TEMPLATE.format(size=len(raw_bytes)))

        if file_name.endswith(XLSX_SUFFIX):
            return self.parse_xlsx_bytes(raw_bytes)

        if file_name.endswith(LEGACY_XLS_SUFFIX):
            self.diagnostics.append(DIAG_UNSUPPORTED_XLS)
            return []

        text = raw_bytes.decode(CSV_TEXT_ENCODING, errors="ignore")
        self.diagnostics.append(
            DIAG_CSV_DECODED_LEN_TEMPLATE.format(length=len(text))
        )

        if text:
            sample = (
                text[:CSV_SAMPLE_LENGTH].replace("\r", "\\r").replace("\n", "\\n")
            )
            self.diagnostics.append(DIAG_CSV_SAMPLE_TEMPLATE.format(sample=sample))

        return self.parse_csv_text(text)

    def parse_csv_text(self, text: str) -> list[SheetMappingRow]:
        if not str(text or "").strip():
            self.diagnostics.append(DIAG_CSV_EMPTY)
            return []

        raw_rows = [tuple(row) for row in csv.reader(StringIO(text))]
        if not raw_rows:
            self.diagnostics.append(DIAG_CSV_NO_ROWS)
            return []

        headers = [self._cell_to_string(value) for value in raw_rows[0]]
        self.diagnostics.append(DIAG_CSV_HEADERS_TEMPLATE.format(headers=headers))

        output = self._parse_data_rows(headers=headers, data_rows=raw_rows[1:])

        self.diagnostics.append(
            DIAG_CSV_USEFUL_ROWS_TEMPLATE.format(count=len(output))
        )
        self._add_population_diagnostics(output)
        return output

    def parse_xlsx_bytes(self, raw_bytes: bytes) -> list[SheetMappingRow]:
        if not raw_bytes:
            self.diagnostics.append(DIAG_XLSX_EMPTY)
            return []

        workbook = load_workbook(
            BytesIO(raw_bytes), data_only=True, read_only=True
        )
        output: list[SheetMappingRow] = []

        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue

            header_index = self._find_header_row(
                rows=rows, sheet_title=worksheet.title
            )
            if header_index < 0:
                self.diagnostics.append(
                    DIAG_SHEET_NO_HEADER_TEMPLATE.format(sheet=worksheet.title)
                )
                continue

            headers = [
                self._cell_to_string(value) for value in rows[header_index]
            ]
            parsed_rows = self._parse_data_rows(
                headers=headers, data_rows=rows[header_index + 1:]
            )
            self.diagnostics.append(
                DIAG_SHEET_USEFUL_ROWS_TEMPLATE.format(
                    sheet=worksheet.title, count=len(parsed_rows)
                )
            )
            output.extend(parsed_rows)

        self.diagnostics.append(
            DIAG_XLSX_USEFUL_ROWS_TEMPLATE.format(count=len(output))
        )
        self._add_population_diagnostics(output)
        return output

    def _parse_data_rows(
        self, *, headers: list[str], data_rows
    ) -> list[SheetMappingRow]:
        output: list[SheetMappingRow] = []
        for row in data_rows:
            raw_row = self._row_to_dict(headers=headers, row=row)
            mapping_row = self._to_mapping_row(raw_row)
            if self._has_useful_content(mapping_row):
                output.append(mapping_row)
        return output